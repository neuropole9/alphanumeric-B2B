from __future__ import annotations

import csv
import json
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
import re
import uuid

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, EmailStr, Field, field_validator
from sqlalchemy import desc, func, or_, select
from sqlalchemy.orm import Session

from . import models as m
from .advanced_services import deliver_email, financial_document_pdf, money, recalc_financial_document, resolve_price, rma_service_pdf
from .config import settings
from .db import get_db
from .deps import ensure_workspace, get_current_user, has_permission, require_csrf
from .operations_api import admin_or, project_access, stock_balance
from .services import audit, next_number

router = APIRouter(prefix="/api/v1")
ADMIN = {"ADMIN", "SUPER_ADMIN"}


def now() -> datetime: return datetime.now(timezone.utc)
def as_float(value): return float(value or 0)


class MoneyLine(BaseModel):
    product_id: str | None = None
    description: str = Field(min_length=1, max_length=500)
    sku: str | None = Field(None, max_length=80)
    hsn_sac: str | None = Field(None, max_length=32)
    quantity: Decimal = Field(gt=0)
    unit: str = Field("Nos", min_length=1, max_length=30)
    unit_price: Decimal = Field(ge=0)
    tax_rate: Decimal = Field(0, ge=0, le=100)


class FinancialDocumentIn(BaseModel):
    document_type: str
    customer_id: str
    project_id: str
    order_id: str | None = None
    invoice_id: str | None = None
    reason: str | None = Field(None, max_length=3000)
    internal_remarks: str | None = Field(None, max_length=3000)
    discount_percent: Decimal = Field(0, ge=0, le=100)
    freight: Decimal = Field(0, ge=0)
    additional_charges: Decimal = Field(0, ge=0)
    items: list[MoneyLine] = Field(default_factory=list)

    @field_validator("document_type")
    @classmethod
    def valid_type(cls, value):
        value = value.upper()
        if value not in {"PROFORMA", "CREDIT_NOTE", "DEBIT_NOTE"}: raise ValueError("Unsupported document type")
        return value


def financial_json(db: Session, row: m.FinancialDocument, detail=False):
    data = {"id": row.id, "number": row.number, "document_type": row.document_type, "customer_id": row.customer_id,
            "project_id": row.project_id, "order_id": row.order_id, "invoice_id": row.invoice_id, "status": row.status,
            "reason": row.reason, "subtotal": as_float(row.subtotal), "discount_amount": as_float(row.discount_amount),
            "tax_total": as_float(row.tax_total), "grand_total": as_float(row.grand_total), "created_at": row.created_at}
    if detail:
        data["items"] = [{"id": x.id, "product_id": x.product_id, "sku": x.sku, "description": x.description,
                          "quantity": as_float(x.quantity), "unit": x.unit, "unit_price": as_float(x.unit_price),
                          "tax_rate": as_float(x.tax_rate), "line_total": as_float(x.line_total)}
                         for x in db.scalars(select(m.FinancialDocumentItem).where(m.FinancialDocumentItem.document_id == row.id)).all()]
    return data


@router.get("/financial-documents")
def financial_documents(document_type: str | None = None, status: str | None = None, page: int = Query(1, ge=1),
                        page_size: int = Query(25, ge=1, le=100), db: Session = Depends(get_db),
                        user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.view")
    stmt = select(m.FinancialDocument)
    if document_type: stmt = stmt.where(m.FinancialDocument.document_type == document_type.upper())
    if status: stmt = stmt.where(m.FinancialDocument.status == status.upper())
    if user.partner_id: stmt = stmt.join(m.Project).where(m.Project.partner_id == user.partner_id)
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    rows = db.scalars(stmt.order_by(desc(m.FinancialDocument.created_at)).offset((page-1)*page_size).limit(page_size)).all()
    return {"items": [financial_json(db, x) for x in rows], "page": page, "page_size": page_size, "total": total}


@router.post("/financial-documents", dependencies=[Depends(require_csrf)])
def create_financial_document(body: FinancialDocumentIn, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.manage")
    project = db.get(m.Project, body.project_id); customer = db.get(m.Customer, body.customer_id)
    if not project or not customer or project.customer_id != customer.id: raise HTTPException(422, "Customer/project mismatch")
    project_access(db, user, project.id, "commercial.manage")
    source_items = list(body.items)
    if body.document_type == "PROFORMA" and body.order_id and not source_items:
        order = db.get(m.Order, body.order_id)
        if not order or order.project_id != project.id: raise HTTPException(422, "Order/project mismatch")
        source_items = [MoneyLine(product_id=x.product_id, description=x.description, sku=x.sku, quantity=x.qty,
                                  unit_price=x.rate, tax_rate=x.tax_rate) for x in db.scalars(select(m.OrderItem).where(m.OrderItem.order_id == order.id)).all()]
    if body.document_type in {"CREDIT_NOTE", "DEBIT_NOTE"}:
        invoice = db.scalar(select(m.Invoice).where(m.Invoice.id == body.invoice_id).with_for_update()) if body.invoice_id else None
        if not invoice or invoice.project_id != project.id: raise HTTPException(422, "An invoice in the project is required")
        if not body.reason: raise HTTPException(422, "A reason is required for an accounting adjustment")
        if not source_items:
            source_items = [MoneyLine(product_id=x.product_id, description=x.description, sku=x.sku, hsn_sac=x.hsn_sac,
                                      quantity=x.qty, unit=x.unit, unit_price=x.rate, tax_rate=x.tax_rate)
                            for x in db.scalars(select(m.InvoiceItem).where(m.InvoiceItem.invoice_id == invoice.id)).all()]
    if not source_items: raise HTTPException(422, "At least one line item is required")
    prefix = {"PROFORMA":"PI", "CREDIT_NOTE":"CN", "DEBIT_NOTE":"DN"}[body.document_type]
    row = m.FinancialDocument(document_type=body.document_type, number=next_number(db, body.document_type.lower(), prefix),
        customer_id=body.customer_id, project_id=body.project_id, order_id=body.order_id, invoice_id=body.invoice_id,
        reason=body.reason, internal_remarks=body.internal_remarks, discount_percent=body.discount_percent,
        freight=body.freight, additional_charges=body.additional_charges, created_by=user.id)
    db.add(row); db.flush()
    for line in source_items:
        db.add(m.FinancialDocumentItem(document_id=row.id, **line.model_dump(), line_subtotal=0, line_tax=0, line_total=0))
    db.flush(); recalc_financial_document(db, row)
    if row.document_type == "CREDIT_NOTE" and row.invoice_id:
        already = db.scalar(select(func.coalesce(func.sum(m.FinancialDocument.grand_total), 0)).where(
            m.FinancialDocument.invoice_id == row.invoice_id, m.FinancialDocument.document_type == "CREDIT_NOTE",
            m.FinancialDocument.status.in_(["DRAFT", "ISSUED"]), m.FinancialDocument.id != row.id)) or 0
        if Decimal(already) + Decimal(row.grand_total) > Decimal(db.get(m.Invoice, row.invoice_id).grand_total):
            raise HTTPException(409, "Credit notes cannot exceed the invoice total")
    audit(db, user.id, "financial_document.created", "financial_document", row.id, {"type": row.document_type})
    db.commit(); return financial_json(db, row, True)


@router.get("/financial-documents/{document_id}")
def financial_document(document_id: str, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.view"); row = db.get(m.FinancialDocument, document_id)
    if not row: raise HTTPException(404, "Financial document not found")
    project_access(db, user, row.project_id); return financial_json(db, row, True)


@router.post("/financial-documents/{document_id}/issue", dependencies=[Depends(require_csrf)])
def issue_financial_document(document_id: str, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.manage"); row = db.get(m.FinancialDocument, document_id)
    if not row: raise HTTPException(404, "Financial document not found")
    if row.status != "DRAFT": raise HTTPException(409, "Only draft documents can be issued")
    row.status, row.issued_by, row.issued_at = "ISSUED", user.id, now(); audit(db, user.id, "financial_document.issued", "financial_document", row.id)
    db.commit(); return financial_json(db, row)


@router.post("/financial-documents/{document_id}/cancel", dependencies=[Depends(require_csrf)])
def cancel_financial_document(document_id: str, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.manage"); row = db.get(m.FinancialDocument, document_id)
    if not row: raise HTTPException(404, "Financial document not found")
    if row.status == "CANCELLED": raise HTTPException(409, "Document is already cancelled")
    row.status, row.cancelled_at = "CANCELLED", now(); audit(db, user.id, "financial_document.cancelled", "financial_document", row.id)
    db.commit(); return financial_json(db, row)


@router.get("/financial-documents/{document_id}/pdf")
def financial_pdf(document_id: str, download: bool = False, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.view"); row = db.get(m.FinancialDocument, document_id)
    if not row: raise HTTPException(404, "Financial document not found")
    project_access(db, user, row.project_id); filename = re.sub(r"[^A-Za-z0-9._-]", "-", row.number) + ".pdf"
    return Response(financial_document_pdf(db, row), media_type="application/pdf", headers={"Content-Disposition": f"{'attachment' if download else 'inline'}; filename=\"{filename}\""})


class PricingRuleIn(BaseModel):
    name: str = Field(min_length=2, max_length=180); scope_type: str = "GENERAL"; product_id: str | None = None
    category_id: str | None = None; customer_id: str | None = None; partner_id: str | None = None; zone: str | None = None
    adjustment_type: str; adjustment_value: Decimal = Field(ge=0); min_quantity: Decimal = Field(0, ge=0)
    max_quantity: Decimal | None = Field(None, ge=0); valid_from: date = Field(default_factory=date.today); valid_until: date | None = None
    priority: int = Field(100, ge=0, le=10000); approval_threshold_percent: Decimal | None = Field(None, ge=0, le=100); status: str = "ACTIVE"


def pricing_json(x: m.PricingRule):
    return {k: getattr(x, k) for k in ("id","name","scope_type","product_id","category_id","customer_id","partner_id","zone",
            "adjustment_type","adjustment_value","min_quantity","max_quantity","valid_from","valid_until","priority","approval_threshold_percent","status")}


@router.get("/pricing-rules")
def pricing_rules(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.view"); return [pricing_json(x) for x in db.scalars(select(m.PricingRule).order_by(desc(m.PricingRule.priority), m.PricingRule.name)).all()]


@router.post("/pricing-rules", dependencies=[Depends(require_csrf)])
def create_pricing_rule(body: PricingRuleIn, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.manage")
    if body.max_quantity is not None and body.max_quantity < body.min_quantity: raise HTTPException(422, "max_quantity must be at least min_quantity")
    if body.valid_until and body.valid_until < body.valid_from: raise HTTPException(422, "valid_until must not precede valid_from")
    if body.adjustment_type not in {"FIXED_PRICE","PERCENT_DISCOUNT","FIXED_DISCOUNT","PERCENT_MARKUP"}: raise HTTPException(422, "Invalid adjustment type")
    row = m.PricingRule(**body.model_dump(), created_by=user.id); db.add(row); db.flush(); audit(db, user.id, "pricing_rule.created", "pricing_rule", row.id); db.commit()
    return pricing_json(row)


@router.patch("/pricing-rules/{rule_id}", dependencies=[Depends(require_csrf)])
def update_pricing_rule(rule_id: str, body: PricingRuleIn, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.manage"); row = db.get(m.PricingRule, rule_id)
    if not row: raise HTTPException(404, "Pricing rule not found")
    for key, value in body.model_dump().items(): setattr(row, key, value)
    audit(db, user.id, "pricing_rule.updated", "pricing_rule", row.id); db.commit(); return pricing_json(row)


@router.get("/pricing/explain")
def pricing_explain(product_id: str, quantity: Decimal = Query(1, gt=0), customer_id: str | None = None,
                    partner_id: str | None = None, zone: str | None = None, db: Session = Depends(get_db),
                    user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.view"); product = db.get(m.Product, product_id)
    if not product: raise HTTPException(404, "Product not found")
    if user.partner_id and partner_id not in {None, user.partner_id}: raise HTTPException(403, "Partner scope denied")
    return resolve_price(db, product, quantity, customer_id=customer_id, partner_id=partner_id or user.partner_id, zone=zone)


class TargetIn(BaseModel):
    name: str = Field(min_length=2, max_length=180); target_type: str; period_type: str = "MONTHLY"
    period_start: date; period_end: date; target_amount: Decimal = Field(ge=0); sales_user_id: str | None = None
    team: str | None = None; zone: str | None = None; partner_id: str | None = None; status: str = "ACTIVE"


def target_details(db: Session, target: m.SalesTarget):
    start, end = target.period_start, target.period_end
    if target.target_type == "REVENUE":
        stmt = select(m.Invoice).where(m.Invoice.invoice_date >= start, m.Invoice.invoice_date <= end, m.Invoice.status != "CANCELLED")
        if target.partner_id: stmt = stmt.join(m.Project, m.Invoice.project_id == m.Project.id).where(m.Project.partner_id == target.partner_id)
        if target.zone: stmt = stmt.join(m.Project, m.Invoice.project_id == m.Project.id).join(m.Partner, m.Project.partner_id == m.Partner.id).where(m.Partner.zone == target.zone)
        rows = db.scalars(stmt).all(); entries = [{"id": x.id, "number": x.number, "date": x.invoice_date, "amount": as_float(x.grand_total)} for x in rows]
    elif target.target_type == "ORDERS":
        rows = db.scalars(select(m.Order).where(m.Order.order_date >= start, m.Order.order_date <= end, m.Order.status != "CANCELLED")).all()
        entries = [{"id": x.id, "number": x.number, "date": x.order_date, "amount": as_float(x.grand_total)} for x in rows]
    else:
        rows = db.scalars(select(m.Inquiry).where(func.date(m.Inquiry.created_at) >= start, func.date(m.Inquiry.created_at) <= end)).all()
        entries = [{"id": x.id, "number": x.number, "date": x.created_at.date(), "amount": 1} for x in rows]
    actual = sum(Decimal(str(x["amount"])) for x in entries); goal = Decimal(target.target_amount or 0)
    return entries, actual, money(actual * 100 / goal) if goal else Decimal(0)


def target_json(db, x, drilldown=False):
    entries, actual, achievement = target_details(db, x)
    data = {k:getattr(x,k) for k in ("id","name","target_type","period_type","period_start","period_end","target_amount","sales_user_id","team","zone","partner_id","status")}
    data.update(actual=actual, achievement_percent=achievement)
    if drilldown: data["entries"] = entries
    return data


@router.get("/sales-targets")
def targets(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "reports.view"); stmt = select(m.SalesTarget)
    if user.partner_id: stmt = stmt.where(m.SalesTarget.partner_id == user.partner_id)
    return [target_json(db, x) for x in db.scalars(stmt.order_by(desc(m.SalesTarget.period_start))).all()]


@router.post("/sales-targets", dependencies=[Depends(require_csrf)])
def create_target(body: TargetIn, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user, "commercial.manage")
    if body.period_end < body.period_start: raise HTTPException(422, "period_end must not precede period_start")
    if body.target_type not in {"REVENUE","ORDERS","INQUIRIES"}: raise HTTPException(422, "Invalid target type")
    row=m.SalesTarget(**body.model_dump(), assigned_by=user.id); db.add(row); db.flush(); audit(db,user.id,"sales_target.created","sales_target",row.id); db.commit()
    return target_json(db,row)


@router.get("/sales-targets/id/{target_id}")
def target_drilldown(target_id: str, db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user,"reports.view"); row=db.get(m.SalesTarget,target_id)
    if not row: raise HTTPException(404,"Sales target not found")
    return target_json(db,row,True)


@router.get("/sales-targets/export.csv")
def target_export(db: Session = Depends(get_db), user: m.User = Depends(get_current_user)):
    admin_or(user,"reports.view"); out=StringIO(); writer=csv.writer(out); writer.writerow(["Name","Type","Period start","Period end","Target","Actual","Achievement %"])
    for row in db.scalars(select(m.SalesTarget).order_by(m.SalesTarget.period_start)).all():
        value=target_json(db,row); writer.writerow([row.name,row.target_type,row.period_start,row.period_end,row.target_amount,value["actual"],value["achievement_percent"]])
    return Response(out.getvalue(),media_type="text/csv",headers={"Content-Disposition":"attachment; filename=sales-targets.csv"})


IMPORT_COLUMNS = ["workspace","category","sku","name","description","brand","unit","price","cost","tax_rate","hsn_sac","reorder_level","lead_time_days","warranty","specs_json"]


def safe_cell(value):
    text = "" if value is None else str(value).strip()
    return "'" + text if text.startswith(("=","+","-","@")) else text


@router.get("/catalogue-import/template.xlsx")
def import_template(user: m.User = Depends(get_current_user)):
    admin_or(user,"commercial.manage"); workbook=Workbook(); sheet=workbook.active; sheet.title="Products"; sheet.append(IMPORT_COLUMNS)
    sheet.append(["LIGHTING","Fixtures","AN-001","Sample fixture","","AlphaNumeric","Nos",1000,700,18,"9405",10,7,"2 years",'{"power_w":"20"}'])
    out=BytesIO(); workbook.save(out); return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=catalogue-import-template.xlsx"})


def read_import(content: bytes, suffix: str):
    if suffix == ".csv":
        return list(csv.DictReader(StringIO(content.decode("utf-8-sig"))))
    book=load_workbook(BytesIO(content),read_only=True,data_only=True); rows=list(book.active.iter_rows(values_only=True));
    if not rows: return []
    headers=[str(x or "").strip() for x in rows[0]]; return [dict(zip(headers,row)) for row in rows[1:] if any(v is not None for v in row)]


def validate_import_row(db,row,seen):
    clean={key:safe_cell(row.get(key)) for key in IMPORT_COLUMNS}; errors=[]; sku=clean["sku"].upper(); clean["sku"]=sku
    for key in ("workspace","category","sku","name","price","tax_rate"): 
        if not clean[key]: errors.append(f"{key} is required")
    if sku in seen: errors.append("duplicate SKU in file")
    seen.add(sku)
    if sku and db.scalar(select(m.Product.id).where(m.Product.sku==sku)): errors.append("SKU already exists")
    for key in ("price","cost","tax_rate","reorder_level"):
        if clean[key]:
            try:
                if Decimal(clean[key]) < 0: errors.append(f"{key} cannot be negative")
            except InvalidOperation: errors.append(f"{key} must be numeric")
    category=db.scalar(select(m.Category).where(func.lower(m.Category.name)==clean["category"].lower(),m.Category.workspace==clean["workspace"].upper())) if clean["category"] and clean["workspace"] else None
    if not category: errors.append("category does not exist in workspace")
    specs={}
    if clean["specs_json"]:
        try: specs=json.loads(clean["specs_json"]); assert isinstance(specs,dict)
        except Exception: errors.append("specs_json must be a JSON object")
    return clean,errors,category,specs


@router.post("/catalogue-imports", dependencies=[Depends(require_csrf)])
async def catalogue_import(file: UploadFile=File(...), dry_run: bool=True, mode: str="ATOMIC", application: str|None=None, db: Session=Depends(get_db), user:m.User=Depends(get_current_user)):
    admin_or(user,"commercial.manage"); mode=mode.upper(); suffix=Path(file.filename or "").suffix.lower()
    locked_workspace=(application or "").upper().strip()
    if locked_workspace not in {"LIGHTING","AUTOMATION"}: raise HTTPException(422,"application must be LIGHTING or AUTOMATION")
    ensure_workspace(user,locked_workspace)
    if suffix not in {".csv",".xlsx"}: raise HTTPException(415,"Only CSV and XLSX files are supported")
    content=await file.read(settings.media_max_bytes+1)
    if len(content)>settings.media_max_bytes: raise HTTPException(413,"Import file is too large")
    if suffix==".xlsx" and not content.startswith(b"PK\x03\x04"): raise HTTPException(422,"Invalid XLSX file signature")
    if suffix==".csv" and b"\x00" in content[:4096]: raise HTTPException(422,"CSV must be plain UTF-8 text")
    if mode not in {"ATOMIC","PARTIAL"}: raise HTTPException(422,"mode must be ATOMIC or PARTIAL")
    try: source=read_import(content,suffix)
    except Exception: raise HTTPException(422,"The import file could not be parsed")
    if len(source)>5000: raise HTTPException(413,"Imports are limited to 5,000 rows")
    job=m.CatalogueImportJob(workspace=locked_workspace,file_name=Path(file.filename or "import").name,mode=mode,dry_run=dry_run,status="PROCESSING",created_by=user.id); db.add(job); db.flush()
    prepared=[]; seen=set()
    for index,raw in enumerate(source,start=2):
        clean,errors,category,specs=validate_import_row(db,raw,seen)
        if clean.get("workspace","").upper()!=locked_workspace: errors.append(f"workspace must be {locked_workspace} for this import")
        prepared.append((index,clean,errors,category,specs))
    any_errors=any(x[2] for x in prepared)
    for index,clean,errors,category,specs in prepared:
        product=None
        may_write=not dry_run and not errors and not(mode=="ATOMIC" and any_errors)
        if may_write:
            product=m.Product(workspace=clean["workspace"].upper(),category_id=category.id,sku=clean["sku"],name=clean["name"],description=clean["description"] or None,
                brand=clean["brand"] or "AlphaNumeric",unit=clean["unit"] or "Nos",price=Decimal(clean["price"]),cost=Decimal(clean["cost"]) if clean["cost"] else None,
                tax_rate=Decimal(clean["tax_rate"]),hsn_sac=clean["hsn_sac"] or None,reorder_level=Decimal(clean["reorder_level"] or 0),
                lead_time_days=int(clean["lead_time_days"]) if clean["lead_time_days"] else None,warranty=clean["warranty"] or None,specs=specs)
            db.add(product); db.flush()
        db.add(m.CatalogueImportRow(job_id=job.id,row_number=index,status="ERROR" if errors else ("VALID" if dry_run or not may_write else "IMPORTED"),
                                    sku=clean["sku"] or None,product_id=product.id if product else None,source_data=clean,errors=errors))
    job.total_rows=len(prepared); job.failed_rows=sum(bool(x[2]) for x in prepared); job.success_rows=job.total_rows-job.failed_rows
    job.status="VALIDATED" if dry_run else ("FAILED" if mode=="ATOMIC" and any_errors else "COMPLETED"); job.completed_at=now()
    audit(db,user.id,"catalogue_import.completed","catalogue_import",job.id,{"dry_run":dry_run,"mode":mode,"failed":job.failed_rows}); db.commit()
    return {"id":job.id,"status":job.status,"dry_run":job.dry_run,"mode":job.mode,"total_rows":job.total_rows,"success_rows":job.success_rows,"failed_rows":job.failed_rows}


@router.get("/catalogue-imports")
def import_history(application:str,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"commercial.view"); ws=application.upper(); ensure_workspace(user,ws)
    return [{"id":x.id,"workspace":x.workspace,"file_name":x.file_name,"mode":x.mode,"dry_run":x.dry_run,"status":x.status,"total_rows":x.total_rows,"success_rows":x.success_rows,"failed_rows":x.failed_rows,"created_at":x.created_at} for x in db.scalars(select(m.CatalogueImportJob).where(m.CatalogueImportJob.workspace==ws).order_by(desc(m.CatalogueImportJob.created_at))).all()]


@router.get("/catalogue-imports/{job_id}/errors.csv")
def import_errors(job_id:str,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"commercial.view"); job=db.get(m.CatalogueImportJob,job_id)
    if not job: raise HTTPException(404,"Import job not found")
    ensure_workspace(user,job.workspace)
    out=StringIO(); writer=csv.writer(out); writer.writerow(["row_number","sku","errors"])
    for row in db.scalars(select(m.CatalogueImportRow).where(m.CatalogueImportRow.job_id==job_id,m.CatalogueImportRow.status=="ERROR")).all(): writer.writerow([row.row_number,safe_cell(row.sku),"; ".join(row.errors)])
    return Response(out.getvalue(),media_type="text/csv",headers={"Content-Disposition":f"attachment; filename=catalogue-import-{job_id}-errors.csv"})


class ApprovalIn(BaseModel):
    entity_type:str; entity_id:str; project_id:str|None=None; rule_code:str=Field(min_length=2,max_length=80); comments:str|None=Field(None,max_length=3000)
class DecisionIn(BaseModel):
    decision:str; comments:str|None=Field(None,max_length=3000)


@router.get("/approval-requests")
def approvals(status:str|None=None,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    stmt=select(m.ApprovalRequest)
    if status: stmt=stmt.where(m.ApprovalRequest.status==status.upper())
    if user.role not in ADMIN and not has_permission(user,"commercial.manage"): stmt=stmt.where(m.ApprovalRequest.requested_by==user.id)
    return [{"id":x.id,"entity_type":x.entity_type,"entity_id":x.entity_id,"project_id":x.project_id,"rule_code":x.rule_code,"status":x.status,"comments":x.comments,"requested_by":x.requested_by,"decided_by":x.decided_by,"created_at":x.created_at,"decided_at":x.decided_at} for x in db.scalars(stmt.order_by(desc(m.ApprovalRequest.created_at))).all()]


@router.post("/approval-requests",dependencies=[Depends(require_csrf)])
def request_approval(body:ApprovalIn,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    project_access(db,user,body.project_id); existing=db.scalar(select(m.ApprovalRequest).where(m.ApprovalRequest.entity_type==body.entity_type,m.ApprovalRequest.entity_id==body.entity_id,m.ApprovalRequest.rule_code==body.rule_code,m.ApprovalRequest.status=="PENDING"))
    if existing: return {"id":existing.id,"status":existing.status,"duplicate":True}
    row=m.ApprovalRequest(**body.model_dump(),requested_by=user.id); db.add(row); db.flush(); audit(db,user.id,"approval.requested","approval",row.id); db.commit(); return {"id":row.id,"status":row.status,"duplicate":False}


@router.post("/approval-requests/{approval_id}/decision",dependencies=[Depends(require_csrf)])
def decide_approval(approval_id:str,body:DecisionIn,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"commercial.manage"); row=db.scalar(select(m.ApprovalRequest).where(m.ApprovalRequest.id==approval_id).with_for_update())
    if not row: raise HTTPException(404,"Approval request not found")
    decision=body.decision.upper()
    if row.status!="PENDING" or decision not in {"APPROVED","REJECTED"}: raise HTTPException(409,"Approval cannot be decided")
    row.status=decision; row.comments=body.comments; row.decided_by=user.id; row.decided_at=now(); audit(db,user.id,"approval.decided","approval",row.id,{"decision":decision}); db.commit(); return {"id":row.id,"status":row.status}


class RMAAssessment(BaseModel):
    warranty_eligible:bool|None=None; warranty_reason:str|None=Field(None,max_length=3000); inspection_notes:str|None=Field(None,max_length=5000)
    diagnosis:str|None=Field(None,max_length=5000); customer_notes:str|None=Field(None,max_length=5000); internal_notes:str|None=Field(None,max_length=5000)
    labour_cost:Decimal=Field(0,ge=0); sla_due_at:datetime|None=None; pickup_reference:str|None=None; return_dispatch_reference:str|None=None
class RMAPartIn(BaseModel):
    product_id:str; warehouse_id:str; quantity:Decimal=Field(gt=0); unit_cost:Decimal=Field(0,ge=0)


@router.patch("/rmas/{rma_id}/assessment",dependencies=[Depends(require_csrf)])
def assess_rma(rma_id:str,body:RMAAssessment,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"rma.manage"); row=db.get(m.RMARequest,rma_id)
    if not row: raise HTTPException(404,"RMA not found")
    for key,value in body.model_dump().items(): setattr(row,key,value)
    audit(db,user.id,"rma.assessed","rma",row.id); db.commit(); return {"id":row.id,"status":row.status,"warranty_eligible":row.warranty_eligible}


@router.post("/rmas/{rma_id}/parts",dependencies=[Depends(require_csrf)])
def add_rma_part(rma_id:str,body:RMAPartIn,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"rma.manage"); rma=db.get(m.RMARequest,rma_id); warehouse=db.scalar(select(m.Warehouse).where(m.Warehouse.id==body.warehouse_id).with_for_update())
    product=db.get(m.Product,body.product_id)
    if not rma or not warehouse or not product: raise HTTPException(404,"RMA, warehouse or product not found")
    before=stock_balance(db,warehouse.id,product.id)
    if before<body.quantity and not warehouse.allow_negative: raise HTTPException(409,f"Insufficient stock; available {before}")
    part=m.RMAPart(rma_id=rma.id,**body.model_dump(),added_by=user.id); db.add(part); db.flush()
    db.add(m.StockLedger(product_id=product.id,warehouse_id=warehouse.id,movement_type="RMA_PART",quantity=body.quantity,direction="OUT",reference_type="RMA",reference_id=rma.id,idempotency_key=f"rma-part:{part.id}",actor_id=user.id))
    audit(db,user.id,"rma.part_added","rma",rma.id,{"product_id":product.id,"quantity":str(body.quantity)}); db.commit(); return {"id":part.id,"balance":as_float(before-body.quantity)}


ALLOWED_EVIDENCE={"image/jpeg","image/png","application/pdf"}
@router.post("/rmas/{rma_id}/evidence",dependencies=[Depends(require_csrf)])
async def upload_rma_evidence(rma_id:str,file:UploadFile=File(...),db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    row=db.get(m.RMARequest,rma_id)
    if not row: raise HTTPException(404,"RMA not found")
    if user.role not in ADMIN and row.requested_by!=user.id and not has_permission(user,"rma.manage"): raise HTTPException(403,"RMA access denied")
    if file.content_type not in ALLOWED_EVIDENCE: raise HTTPException(415,"Only PDF, PNG and JPEG evidence is supported")
    content=await file.read(settings.media_max_bytes+1)
    if len(content)>settings.media_max_bytes: raise HTTPException(413,"Evidence file is too large")
    signatures={"image/jpeg":(b"\xff\xd8\xff",),"image/png":(b"\x89PNG\r\n\x1a\n",),"application/pdf":(b"%PDF-",)}
    if not any(content.startswith(signature) for signature in signatures[file.content_type]): raise HTTPException(422,"File content does not match its declared type")
    extension={"image/jpeg":".jpg","image/png":".png","application/pdf":".pdf"}[file.content_type]; key=f"rma/{rma_id}/{uuid.uuid4().hex}{extension}"
    target=(Path(settings.media_root).resolve()/key); target.parent.mkdir(parents=True,exist_ok=True); target.write_bytes(content)
    evidence=m.RMAEvidence(rma_id=rma_id,file_name=Path(file.filename or "evidence").name,storage_key=key,mime_type=file.content_type,size_bytes=len(content),uploaded_by=user.id); db.add(evidence); db.flush(); audit(db,user.id,"rma.evidence_uploaded","rma",rma_id,{"evidence_id":evidence.id}); db.commit()
    return {"id":evidence.id,"file_name":evidence.file_name,"mime_type":evidence.mime_type,"size_bytes":evidence.size_bytes}


@router.get("/rmas/{rma_id}/evidence")
def list_rma_evidence(rma_id:str,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    row=db.get(m.RMARequest,rma_id)
    if not row: raise HTTPException(404,"RMA not found")
    if user.role not in ADMIN and row.requested_by!=user.id and not has_permission(user,"rma.view"): raise HTTPException(403,"RMA access denied")
    return [{"id":x.id,"file_name":x.file_name,"mime_type":x.mime_type,"size_bytes":x.size_bytes,"created_at":x.created_at} for x in db.scalars(select(m.RMAEvidence).where(m.RMAEvidence.rma_id==rma_id)).all()]


@router.get("/rmas/{rma_id}/service-report.pdf")
def rma_report(rma_id:str,download:bool=False,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    row=db.get(m.RMARequest,rma_id)
    if not row:raise HTTPException(404,"RMA not found")
    if user.role not in ADMIN and row.requested_by!=user.id and not has_permission(user,"rma.view"):raise HTTPException(403,"RMA access denied")
    return Response(rma_service_pdf(db,row),media_type="application/pdf",headers={"Content-Disposition":f"{'attachment' if download else 'inline'}; filename=\"{row.number}-service-report.pdf\""})


@router.get("/rmas/{rma_id}/evidence/{evidence_id}/download")
def download_rma_evidence(rma_id:str,evidence_id:str,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    row=db.get(m.RMARequest,rma_id); evidence=db.scalar(select(m.RMAEvidence).where(m.RMAEvidence.id==evidence_id,m.RMAEvidence.rma_id==rma_id))
    if not row or not evidence:raise HTTPException(404,"RMA evidence not found")
    if user.role not in ADMIN and row.requested_by!=user.id and not has_permission(user,"rma.view"):raise HTTPException(403,"RMA access denied")
    root=Path(settings.media_root).resolve(); target=(root/evidence.storage_key).resolve()
    if root not in target.parents or not target.is_file():raise HTTPException(404,"Evidence file is unavailable")
    return FileResponse(target,media_type=evidence.mime_type,filename=evidence.file_name)


@router.get("/rmas/{rma_id}/service-detail")
def rma_service_detail(rma_id:str,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    row=db.get(m.RMARequest,rma_id)
    if not row:raise HTTPException(404,"RMA not found")
    if user.role not in ADMIN and row.requested_by!=user.id and not has_permission(user,"rma.view"):raise HTTPException(403,"RMA access denied")
    history=db.scalars(select(m.RMAHistory).where(m.RMAHistory.rma_id==rma_id).order_by(m.RMAHistory.created_at)).all()
    parts=db.scalars(select(m.RMAPart).where(m.RMAPart.rma_id==rma_id)).all()
    evidence=db.scalars(select(m.RMAEvidence).where(m.RMAEvidence.rma_id==rma_id)).all()
    return {"id":row.id,"number":row.number,"status":row.status,"warranty_eligible":row.warranty_eligible,"warranty_reason":row.warranty_reason,
            "inspection_notes":row.inspection_notes,"diagnosis":row.diagnosis,"customer_notes":row.customer_notes,"internal_notes":row.internal_notes,
            "labour_cost":row.labour_cost,"sla_due_at":row.sla_due_at,"pickup_reference":row.pickup_reference,"return_dispatch_reference":row.return_dispatch_reference,
            "history":[{"from":x.from_status,"to":x.to_status,"notes":x.notes,"created_at":x.created_at} for x in history],
            "parts":[{"id":x.id,"product_id":x.product_id,"warehouse_id":x.warehouse_id,"quantity":x.quantity,"unit_cost":x.unit_cost} for x in parts],
            "evidence":[{"id":x.id,"file_name":x.file_name,"mime_type":x.mime_type,"size_bytes":x.size_bytes} for x in evidence]}


class EmailIn(BaseModel):
    template_key:str=Field(min_length=2,max_length=80); recipient:EmailStr; subject:str=Field(min_length=1,max_length=255)
    text_body:str=Field(min_length=1,max_length=20000); html_body:str=Field(min_length=1,max_length=30000); idempotency_key:str=Field(min_length=4,max_length=180)
    related_type:str|None=None; related_id:str|None=None


def email_json(x): return {"id":x.id,"template_key":x.template_key,"recipient":x.recipient,"subject":x.subject,"status":x.status,"attempts":x.attempts,"error_message":x.error_message,"created_at":x.created_at,"sent_at":x.sent_at}


@router.get("/email-deliveries")
def email_deliveries(db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"reports.view"); return [email_json(x) for x in db.scalars(select(m.EmailDelivery).order_by(desc(m.EmailDelivery.created_at)).limit(250)).all()]


@router.post("/email-deliveries",dependencies=[Depends(require_csrf)])
def queue_email(body:EmailIn,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"commercial.manage"); existing=db.scalar(select(m.EmailDelivery).where(m.EmailDelivery.idempotency_key==body.idempotency_key))
    if existing:return email_json(existing)
    row=m.EmailDelivery(**body.model_dump());db.add(row);db.flush();deliver_email(row);audit(db,user.id,"email.delivery_attempted","email_delivery",row.id,{"status":row.status});db.commit();return email_json(row)


@router.post("/email-deliveries/{delivery_id}/retry",dependencies=[Depends(require_csrf)])
def retry_email(delivery_id:str,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"commercial.manage");row=db.get(m.EmailDelivery,delivery_id)
    if not row:raise HTTPException(404,"Email delivery not found")
    if row.status=="SENT":raise HTTPException(409,"Email was already sent")
    deliver_email(row);audit(db,user.id,"email.delivery_retried","email_delivery",row.id,{"status":row.status});db.commit();return email_json(row)


class PreferenceIn(BaseModel):
    event_key:str=Field(min_length=2,max_length=80);in_app_enabled:bool=True;email_enabled:bool=True
@router.get("/notification-preferences")
def preferences(db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    return [{"event_key":x.event_key,"in_app_enabled":x.in_app_enabled,"email_enabled":x.email_enabled} for x in db.scalars(select(m.NotificationPreference).where(m.NotificationPreference.user_id==user.id)).all()]
@router.put("/notification-preferences/{event_key}",dependencies=[Depends(require_csrf)])
def set_preference(event_key:str,body:PreferenceIn,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    if event_key!=body.event_key:raise HTTPException(422,"Event key mismatch")
    row=db.scalar(select(m.NotificationPreference).where(m.NotificationPreference.user_id==user.id,m.NotificationPreference.event_key==event_key))
    if not row:row=m.NotificationPreference(user_id=user.id,event_key=event_key);db.add(row)
    row.in_app_enabled=body.in_app_enabled;row.email_enabled=body.email_enabled;db.commit();return {"event_key":row.event_key,"in_app_enabled":row.in_app_enabled,"email_enabled":row.email_enabled}


class SerialIn(BaseModel):
    serial_number:str=Field(min_length=2,max_length=120);product_id:str;batch_number:str|None=None;warehouse_id:str|None=None
    customer_id:str|None=None;project_id:str|None=None;room_id:str|None=None;dispatch_id:str|None=None
    warranty_start:date|None=None;warranty_end:date|None=None;status:str="IN_STOCK"


@router.get("/serial-units")
def serial_units(q:str|None=None,status:str|None=None,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"inventory.view");stmt=select(m.SerialUnit)
    if q:stmt=stmt.where(m.SerialUnit.serial_number.ilike(f"%{q}%"))
    if status:stmt=stmt.where(m.SerialUnit.status==status.upper())
    return [{k:getattr(x,k) for k in ("serial_number","product_id","batch_number","warehouse_id","status","customer_id","project_id","room_id","dispatch_id","warranty_start","warranty_end")} for x in db.scalars(stmt.order_by(m.SerialUnit.serial_number).limit(500)).all()]


@router.post("/serial-units",dependencies=[Depends(require_csrf)])
def create_serial(body:SerialIn,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"inventory.receive")
    if db.get(m.SerialUnit,body.serial_number):raise HTTPException(409,"Serial number already exists")
    if not db.get(m.Product,body.product_id):raise HTTPException(404,"Product not found")
    if body.warranty_end and body.warranty_start and body.warranty_end<body.warranty_start:raise HTTPException(422,"Warranty end must not precede start")
    row=m.SerialUnit(**body.model_dump());db.add(row);audit(db,user.id,"serial.created","serial",row.serial_number);db.commit();return {"serial_number":row.serial_number,"status":row.status}


class EntityStatus(BaseModel):
    status:str;reason:str|None=Field(None,max_length=1000);resolution:str|None=Field(None,max_length=5000)


@router.post("/tickets/{ticket_id}/status",dependencies=[Depends(require_csrf)])
def ticket_status(ticket_id:str,body:EntityStatus,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"tickets.manage");row=db.get(m.SupportTicket,ticket_id)
    if not row:raise HTTPException(404,"Ticket not found")
    target=body.status.upper();allowed={"OPEN":{"IN_PROGRESS","CLOSED"},"IN_PROGRESS":{"RESOLVED","CLOSED"},"RESOLVED":{"CLOSED","IN_PROGRESS"},"CLOSED":set()}
    if target not in allowed.get(row.status,set()):raise HTTPException(409,f"Invalid ticket transition {row.status} -> {target}")
    old=row.status;row.status=target;row.resolution=body.resolution or row.resolution;audit(db,user.id,"ticket.status_changed","ticket",row.id,{"from":old,"to":target});db.commit();return {"id":row.id,"status":row.status,"resolution":row.resolution}


@router.post("/warehouses/{warehouse_id}/status",dependencies=[Depends(require_csrf)])
def warehouse_status(warehouse_id:str,body:EntityStatus,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"inventory.receive");row=db.get(m.Warehouse,warehouse_id)
    if not row:raise HTTPException(404,"Warehouse not found")
    target=body.status.upper()
    if target not in {"ACTIVE","INACTIVE"}:raise HTTPException(422,"Status must be ACTIVE or INACTIVE")
    row.status=target;audit(db,user.id,"warehouse.status_changed","warehouse",row.id,{"status":target,"reason":body.reason});db.commit();return {"id":row.id,"status":row.status}


@router.post("/announcements/{announcement_id}/status",dependencies=[Depends(require_csrf)])
def announcement_status(announcement_id:str,body:EntityStatus,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"announcements.manage");row=db.get(m.Announcement,announcement_id)
    if not row:raise HTTPException(404,"Announcement not found")
    target=body.status.upper()
    if target not in {"DRAFT","PUBLISHED","ARCHIVED"}:raise HTTPException(422,"Invalid announcement status")
    row.status=target;audit(db,user.id,"announcement.status_changed","announcement",row.id,{"status":target});db.commit();return {"id":row.id,"status":row.status}


@router.post("/resources/{resource_id}/status",dependencies=[Depends(require_csrf)])
def resource_status(resource_id:str,body:EntityStatus,db:Session=Depends(get_db),user:m.User=Depends(get_current_user)):
    admin_or(user,"resources.manage");row=db.get(m.ResourceAsset,resource_id)
    if not row:raise HTTPException(404,"Resource not found")
    target=body.status.upper()
    if target not in {"PUBLISHED","ARCHIVED"}:raise HTTPException(422,"Invalid resource status")
    row.status=target;audit(db,user.id,"resource.status_changed","resource",row.id,{"status":target});db.commit();return {"id":row.id,"status":row.status}
